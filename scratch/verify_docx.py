from docx import Document

doc = Document("SmartPCB-XAI_Research_Paper_8Pages.docx")
print("Total Paragraphs:", len(doc.paragraphs))
print("Total Tables:", len(doc.tables))

page_breaks = 0
for p in doc.paragraphs:
    for r in p.runs:
        if "w:br" in r._r.xml and "page" in r._r.xml:
            page_breaks += 1

print(f"Total Page Breaks: {page_breaks}")
print(f"Total Pages (Breaks + 1): {page_breaks + 1}")
